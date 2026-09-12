"""Excel report generation (spec section 9) using pandas + openpyxl."""
from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import settings
from database.supabase_client import db

COLUMNS = ["Date", "Time", "Register No", "Student Name", "Class", "Department", "Subject", "Status"]
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _join(rows: list[dict[str, Any]]) -> pd.DataFrame:
    """Join attendance rows with student/subject names into report columns."""
    students = {s["id"]: s for s in db.table("students").select("*").execute().data}
    subjects = {s["id"]: s for s in db.table("subjects").select("*").execute().data}
    records = []
    for r in rows:
        stu = students.get(r["student_id"], {})
        sub = subjects.get(r.get("subject_id"), {})
        records.append({
            "Date": r.get("date"),
            "Time": str(r.get("entry_time") or "")[:5],
            "Register No": stu.get("register_no", ""),
            "Student Name": stu.get("name", ""),
            "Class": stu.get("class", ""),
            "Department": stu.get("department", ""),
            "Subject": sub.get("code", ""),
            "Status": r.get("status", ""),
        })
    return pd.DataFrame(records, columns=COLUMNS)


def _styled_path(df: pd.DataFrame, path: Path, sheet: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet)
        ws = writer.sheets[sheet]
        for cell in ws[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center")
        for idx, col in enumerate(df.columns, start=1):
            if not df.empty:
                width = max(len(str(col)) + 4, *(len(str(v)) + 2 for v in df[col]))
            else:
                width = len(str(col)) + 4
            ws.column_dimensions[get_column_letter(idx)].width = width
        ws.freeze_panes = "A2"
    return path


def export_daily(day: None | str = None) -> Path:
    """Attendance_<YYYY-MM-DD>.xlsx for one day."""
    day = day or date.today().isoformat()
    rows = db.table("attendance").select("*").eq("date", day).execute().data
    df = _join(rows)
    return _styled_path(df, settings.REPORTS_DIR / f"Attendance_{day}.xlsx", "Attendance")


def export_range(start: str, end: str) -> Path:
    """Attendance_<start>_to_<end>.xlsx for a date range."""
    rows = (
        db.table("attendance")
        .select("*")
        .gte("date", start)
        .lte("date", end)
        .execute()
        .data
    )
    df = _join(rows)
    return _styled_path(df, settings.REPORTS_DIR / f"Attendance_{start}_to_{end}.xlsx", "Attendance")


def export_exam_attendance(exam: dict[str, Any]) -> Path:
    """Exam attendance with hall/seat columns."""
    allocations = {
        a["student_id"]: a
        for a in db.table("seat_allocations").select("*").eq("exam_id", exam["id"]).execute().data
    }
    halls = {h["id"]: h for h in db.table("halls").select("*").execute().data}
    seats = {s["id"]: s for s in db.table("seats").select("*").execute().data}
    attendance = {
        r["student_id"]: r
        for r in db.table("exam_attendance").select("*").eq("exam_id", exam["id"]).execute().data
    }
    students = {s["id"]: s for s in db.table("students").select("*").execute().data}

    records = []
    for sid, alloc in allocations.items():
        stu = students.get(sid, {})
        hall = halls.get(alloc["hall_id"], {})
        seat = seats.get(alloc["seat_id"], {})
        att = attendance.get(sid)
        records.append({
            "Exam": exam.get("exam_name", ""),
            "Date": exam.get("exam_date", ""),
            "Register No": stu.get("register_no", ""),
            "Student Name": stu.get("name", ""),
            "Hall": hall.get("hall_name", ""),
            "Seat No": seat.get("seat_number", ""),
            "Status": att.get("status", "") if att else "Absent",
            "Verification Time": str(att.get("verification_time", ""))[:19].replace("T", " ") if att else "",
        })
    df = pd.DataFrame(records, columns=list(records[0].keys()) if records else
                      ["Exam", "Date", "Register No", "Student Name", "Hall", "Seat No", "Status", "Verification Time"])
    safe_name = str(exam.get("exam_name", "exam")).replace(" ", "_")
    return _styled_path(df, settings.REPORTS_DIR / f"Exam_{safe_name}.xlsx", "Exam Attendance")
